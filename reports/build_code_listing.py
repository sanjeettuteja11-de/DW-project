"""Complete code listing of the project as a plain Word document.

    python reports/build_code_listing.py   ->  submission/01b_Complete_Code.docx
"""
import importlib.util
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT  # noqa: F401
from docx.oxml.ns import qn
from docx.shared import Pt, Inches, RGBColor

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "notebooks"))
spec = importlib.util.spec_from_file_location("nbb", ROOT / "notebooks" / "build_submission_notebook.py")
src = (ROOT / "notebooks" / "build_submission_notebook.py").read_text(encoding="utf-8")
ns = {}
exec(src.split("def file_cell")[0].split("REPO = HERE.parent", 1)[1].join(["", ""]) if False else "", ns)
# reuse the file list from the notebook builder
start = src.index("CODE_FILES = [")
end = src.index("def file_cell")
exec(src[start:end], ns)
CODE_FILES, TEST_FILES = ns["CODE_FILES"], ns["TEST_FILES"]
pb = importlib.util.spec_from_file_location("build_pbip", ROOT / "dashboard" / "build_pbip.py")
b = importlib.util.module_from_spec(pb)
pb.loader.exec_module(b)

doc = Document()
sec = doc.sections[0]
sec.page_height, sec.page_width = Inches(11.69), Inches(8.27)
for side in ("left_margin", "right_margin"):
    setattr(sec, side, Inches(0.7))
sec.top_margin = sec.bottom_margin = Inches(0.7)


def font(style, name, size, bold=None):
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is not None:
        for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            rf.attrib.pop(qn(a), None)
    ppr = style.element.pPr
    if ppr is not None and ppr.find(qn("w:pBdr")) is not None:
        ppr.remove(ppr.find(qn("w:pBdr")))
    style.font.name = name
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        style.font.bold = bold
    style.element.rPr.rFonts.set(qn("w:eastAsia"), name)


font(doc.styles["Normal"], "Times New Roman", 11)
for lvl, size in ((1, 15), (2, 12.5)):
    st = doc.styles[f"Heading {lvl}"]
    font(st, "Times New Roman", size, True)
    st.paragraph_format.space_before = Pt(12)
    st.paragraph_format.space_after = Pt(4)
doc.styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
font(doc.styles["Title"], "Times New Roman", 20, True)


def code_block(text):
    for i, line in enumerate(text.splitlines(), 1):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(f"{i:>4}  {line}")
        r.font.name = "Courier New"
        r.font.size = Pt(7.5)
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Courier New")


def lines(rel):
    return len((ROOT / rel).read_text(encoding="utf-8").splitlines())


doc.add_paragraph("IPL Phase-Performance and Style-Drift Data Warehouse: Complete Code", style="Title")
doc.add_paragraph("Data Wrangling group project · B.Tech CSE (Data Science), NMIMS MPSTME · Semester V (2026-27)")
doc.add_paragraph("Group members: Sanjeet Singh Tuteja (L076) · [Member 2 name, roll no.] · [Member 3 name, roll no.]")

doc.add_heading("How the code is organised and run", 1)
total = sum(lines(f) for f, _, _ in CODE_FILES if f.endswith((".py", ".sql")))
doc.add_paragraph(
    f"The project is a Python package (src/) with two SQL scripts (sql/) and four mapping tables (mappings/): "
    f"{total:,} lines of Python and SQL in all. The files are listed below in the order the pipeline uses them "
    "(steps E1–E12 are the extraction, cleaning, transformation, loading, validation and export steps). The "
    "same code runs from the command line, from the notebook (01_IPL_DW_Project_Notebook.ipynb, which contains "
    "every file below in section 0 and runs it with outputs) and under the automated tests (Appendix A).")
for cmd in ("pip install -r requirements.txt",
            "python -m src.pipeline --source-url <Kaggle URL> --dataset-version 3   (steps E1–E12)",
            "python -m pytest   (38 tests)"):
    p = doc.add_paragraph(cmd, style="List Bullet")
    for r in p.runs:
        r.font.name = "Courier New"
        r.font.size = Pt(9.5)

rows = [(i, f, step, lines(f), what.replace("`", "")) for i, (f, step, what) in enumerate(CODE_FILES, 1)]
t = doc.add_table(rows=1, cols=5)
t.style = "Table Grid"
for c, h in zip(t.rows[0].cells, ("No.", "File", "Step", "Lines", "What it does")):
    c.text = h
    c.paragraphs[0].runs[0].bold = True
for r in rows:
    cells = t.add_row().cells
    for c, v in zip(cells, r):
        c.text = str(v)
widths = (Inches(0.45), Inches(1.75), Inches(0.75), Inches(0.5), Inches(3.4))
t.autofit = False
for col, w in zip(t.columns, widths):
    col.width = w
for row in t.rows:
    for c, w in zip(row.cells, widths):
        c.width = w
        for p in c.paragraphs:
            for run in p.runs:
                run.font.size = Pt(9)

for i, (f, step, what) in enumerate(CODE_FILES, 1):
    doc.add_heading(f"{i}. {f} ({step})", 1 if i == 1 else 1)
    doc.add_paragraph(what.replace("`", ""))
    code_block((ROOT / f).read_text(encoding="utf-8") or "(empty file)")

doc.add_heading("Appendix A. Automated tests", 1)
doc.add_paragraph("Thirty-eight pytest tests: cleaning rules, SCD Type 2 logic, every validation gate failing on "
                  "deliberately broken data, the full pipeline on a small synthetic dataset, and the analysis, model "
                  "and dashboard code. All 38 pass.")
for f in TEST_FILES:
    doc.add_heading(f, 2)
    code_block((ROOT / f).read_text(encoding="utf-8") or "(empty file: marks tests/ as a package)")

doc.add_heading("Appendix B. Power BI model: relationships and DAX measures", 1)
doc.add_paragraph("The Power BI project (03c_PowerBI_Dashboard_IPL_Scorebook.zip) loads the 18 tables written by "
                  "src/powerbi.py (the Model table is called MLModel there). Relationships run from the many side to "
                  "the one side; inactive ones avoid ambiguous filter paths.")
code_block("\n".join(f"{ft}[{fc}]  *-->1  {tt}[{tc}]" + ("" if a else "   (inactive)")
                     for ft, fc, tt, tc, a in b.RELATIONSHIPS))
doc.add_heading(f"DAX measures ({len(b.MEASURES)})", 2)
dax = []
for table in b.TABLE_ORDER:
    ms = [m for m in b.MEASURES if m[0] == table]
    if ms:
        dax.append(f"-- table {table}")
        dax += [f"{n} = {e}" + (f"    -- format {fm}" if fm else "") for _, n, e, fm, _ in ms]
        dax.append("")
code_block("\n".join(dax).rstrip())

out = ROOT / "submission" / "01b_Complete_Code.docx"
doc.save(out)
print("wrote", out)
